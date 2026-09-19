from __future__ import annotations

import csv
import io
import json
import tempfile
import zipfile
from pathlib import Path
from typing import Any

import geopandas as gpd
from openpyxl import Workbook
from openpyxl.styles import Font

from firewatch.services import pdf_report

VECTOR_FORMATS = {
    "geojson": ("application/geo+json", "geojson"),
    "shp": ("application/zip", "zip"),
    "gpkg": ("application/geopackage+sqlite3", "gpkg"),
    "kml": ("application/vnd.google-earth.kml+xml", "kml"),
}
REPORT_FORMATS = {
    "json": ("application/json", "json"),
    "csv": ("text/csv; charset=utf-8", "csv"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "pdf": ("application/pdf", "pdf"),
}
SHAPEFILE_FIELDS = {"fire_event_id": "event_id", "severity_label": "sev_label", "polygon_id": "poly_id"}


class ExportUnavailableError(RuntimeError):
    pass


def export_vector(collection: dict[str, Any], fmt: str, layer: str) -> bytes:
    if fmt == "geojson":
        crs = {"type": "name", "properties": {"name": "urn:ogc:def:crs:OGC:1.3:CRS84"}}
        return json.dumps({**collection, "name": layer, "crs": crs}, ensure_ascii=False).encode()
    if not collection["features"]:
        raise ExportUnavailableError("В запросе нет контуров: нечего выгружать в этот формат")
    frame = gpd.GeoDataFrame.from_features(collection["features"], crs="EPSG:4326")
    with tempfile.TemporaryDirectory(prefix="fw_export_") as tmp:
        folder = Path(tmp)
        if fmt == "shp":
            frame.rename(columns=SHAPEFILE_FIELDS).to_file(
                folder / f"{layer}.shp", driver="ESRI Shapefile", encoding="UTF-8"
            )
            return _zip_folder(folder)
        driver = {"gpkg": "GPKG", "kml": "KML"}[fmt]
        target = folder / f"{layer}.{fmt}"
        frame.to_file(target, driver=driver, layer=layer)
        return target.read_bytes()


def _zip_folder(folder: Path) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(folder.iterdir()):
            archive.write(path, path.name)
    return buffer.getvalue()


def _report_rows(report: dict[str, Any]) -> list[tuple[str, str, Any, str]]:
    burn, hot, plan = report["burn"], report["hotspots"], report["restoration"]
    power, emissions = hot["power"], report["emissions"]
    place = report["territory"]
    # Территория идёт первой во всех форматах: без неё таблица показателей не
    # отличима от такой же таблицы по соседнему региону.
    rows: list[tuple[str, str, Any, str]] = [
        ("Территория", "Название", place["name"], ""),
        ("Территория", "Тип запроса", place["kind_label"], ""),
        ("Территория", "Площадь запроса", place["area_km2"], "км²"),
        ("Территория", "Координаты", place["bbox_label"], ""),
        ("Территория", "Система координат", place["crs"], ""),
        ("Территория", "Проекция расчёта площадей", place["area_crs"], ""),
        ("Территория", "Покрыто территорией мониторинга", place["monitoring_area"]["covered_km2"], "км²"),
    ]
    rows += [
        ("Субъекты РФ", s["name"], s["area_km2"], "км²") for s in place["subjects"]
    ]
    rows += [
        ("Запрос", "Период с", report["query"]["date_from"], ""),
        ("Запрос", "Период по", report["query"]["date_to"], ""),
        ("Гари", "Площадь всего", burn["total_ha"], "га"),
        ("Гари", "Сцен в расчёте", burn["scenes"], "шт"),
    ]
    rows += [("Гари по степеням", f"{s['label']} степень", s["area_ha"], "га") for s in burn["by_severity"]]
    rows += [
        ("Гари по степеням", f"{s['label']} степень, доля", round(s["share"] * 100, 1), "%")
        for s in burn["by_severity"]
    ]
    rows += [("Гари по покрову", c["label"], c["area_ha"], "га") for c in burn["by_landcover"]]
    rows += [
        ("Термоточки", "Термоаномалий всего", hot["raw_total"], "шт"),
        ("Термоточки", "Природные пожары", hot["fire"], "шт"),
        ("Термоточки", "Очагов природных пожаров", hot["fire_clusters"], "шт"),
        ("Термоточки", "Техногенные источники", hot["technogenic"], "шт"),
        ("Термоточки", "Блики", hot["glint"], "шт"),
        ("Термоточки", "Ложных выездов предотвращено", hot["false_dispatches_prevented"], "шт"),
    ]
    rows += [
        ("Термоточки", "Суммарная мощность излучения", power["total_mw"], "МВт"),
        ("Термоточки", "Самая мощная термоточка", power["max_mw"], "МВт"),
        ("Термоточки", "Насыщенных пикселов I4", power["saturated"], "шт"),
    ]
    rows += [("Выбросы", f"{s['label']}, {s['name']}", s["tonnes"], "т") for s in emissions["species"]]
    rows += [
        ("Выбросы", "Сгоревшая биомасса", emissions["dry_matter_t"], "т"),
        ("Выбросы", "CO₂-эквивалент", emissions["co2e_t"], "т"),
    ]
    rows += [("Восстановление", i["label"], i["area_ha"], "га") for i in plan["items"]]
    rows += [
        ("Восстановление", f"{i['label']}, затраты", i["cost"], plan["currency"])
        for i in plan["items"]
        if i["cost"]
    ]
    rows.append(("Восстановление", "Площадь в расчёте ущерба", plan["damage_area_ha"], "га"))
    return rows


def export_report(report: dict[str, Any], fmt: str) -> bytes:
    if fmt == "json":
        return json.dumps(report, ensure_ascii=False, indent=2).encode()
    if fmt == "pdf":
        try:
            return pdf_report.render(report)
        except pdf_report.PdfUnavailableError as exc:
            raise ExportUnavailableError(str(exc)) from exc
    rows = _report_rows(report)
    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(("section", "metric", "value", "unit"))
        writer.writerows(rows)
        return ("\ufeff" + buffer.getvalue()).encode()
    return _xlsx(report, rows)


def _xlsx(report: dict[str, Any], rows: list[tuple[str, str, Any, str]]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Справка"
    sheet.append(("Раздел", "Показатель", "Значение", "Ед."))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 58
    sheet.column_dimensions["C"].width = 16

    months = book.create_sheet("По месяцам")
    months.append(("Месяц", "Термоточки пожаров", "Площадь гарей, га", "Контуров"))
    burn_months = {m["month"]: m for m in report["burn"]["by_month"]}
    fire_months = {m["month"]: m["n"] for m in report["hotspots"]["by_month"]}
    for month in sorted(burn_months | fire_months):
        burn = burn_months.get(month, {"area_ha": 0.0, "contours": 0})
        months.append((month, fire_months.get(month, 0), burn["area_ha"], burn["contours"]))

    cover = book.create_sheet("По типу территории")
    cover.append(("Тип территории", "Площадь гарей, га", "Термоточки пожаров"))
    fire_cover = {r["label"]: r["n"] for r in report["hotspots"]["by_landcover"]}
    burn_cover = {r["label"]: r["area_ha"] for r in report["burn"]["by_landcover"]}
    for label in sorted(burn_cover | fire_cover):
        cover.append((label, burn_cover.get(label, 0.0), fire_cover.get(label, 0)))

    days = book.create_sheet("Термоточки по дням")
    days.append(("Дата", "Природные", "Техногенные", "Блики"))
    for day in report["hotspots"]["by_day"]:
        days.append((day["day"], day["fire"], day["technogenic"], day["glint"]))

    method = book.create_sheet("Методика")
    for title, body in report["method"].items():
        method.append((title, body))
    method.append(("geometry", json.dumps(report["query"]["geometry"])[:32000]))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
